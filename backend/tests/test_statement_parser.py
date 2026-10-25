import csv
import io
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path
from uuid import uuid4
from zipfile import ZIP_DEFLATED, ZipFile
from zoneinfo import ZoneInfo

import pytest
from openpyxl import Workbook

from src.db.models import BankSource, OperationType
from src.modules.statements.parser import MAX_FILE_BYTES, StatementParser, clean_merchant
from src.modules.statements.schemas import StatementError

FIXTURES = Path(__file__).parent / "fixtures"


@pytest.mark.parametrize(
    ("bank", "count", "skipped"),
    [(BankSource.TBANK, 5, 2), (BankSource.SBER, 3, 0), (BankSource.VTB, 2, 1)],
)
@pytest.mark.parametrize("encoding", ["utf-8-sig", "cp1251"])
def test_supported_csv_formats(bank: BankSource, count: int, skipped: int, encoding: str) -> None:
    name = {BankSource.TBANK: "tbank", BankSource.SBER: "sber", BankSource.VTB: "vtb"}[bank]
    source = (FIXTURES / f"{name}.csv").read_text()
    if encoding == "cp1251":
        source = source.replace("₽", "руб.")
    parsed = StatementParser(ZoneInfo("Europe/Moscow")).parse(
        content=source.encode(encoding),
        filename="statement.csv",
        user_id=uuid4(),
        bank_hint=None,
    )
    assert parsed.bank_source is bank
    assert len(parsed.transactions) == count
    assert parsed.skipped_rows == skipped
    assert all(row.amount > 0 and row.transaction_date.tzinfo is UTC for row in parsed.transactions)
    assert all(row.currency == "RUB" for row in parsed.transactions)
    if bank is BankSource.SBER:
        item = parsed.transactions[0]
        assert item.amount == Decimal("1250.50")
        assert item.operation_type is OperationType.DEBIT
        assert item.transaction_date == datetime(2026, 10, 2, 5, 30, tzinfo=UTC)
        assert item.raw_description == "PYATEROCHKA 12344"
        assert item.merchant_name_clean == "pyaterochka"
        assert item.mcc_code is None
        assert item.bank_category_raw == "Продукты"


@pytest.mark.parametrize("bank", ["sber", "vtb"])
def test_xlsx_and_csv_have_identical_normalization(bank: str) -> None:
    content = (FIXTURES / f"{bank}.csv").read_bytes()
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    for row in csv.reader(io.StringIO(content.decode()), delimiter=";"):
        sheet.append(row)
    if bank == "sber":
        sheet.cell(2, 1).value = datetime(2026, 10, 2, tzinfo=UTC).replace(tzinfo=None)
    binary = io.BytesIO()
    workbook.save(binary)
    parser = StatementParser(ZoneInfo("Europe/Moscow"))
    user_id = uuid4()
    csv_result = parser.parse(content=content, filename="a.csv", user_id=user_id, bank_hint=None)
    xlsx_result = parser.parse(
        content=binary.getvalue(), filename="a.xlsx", user_id=user_id, bank_hint=None
    )
    assert csv_result == xlsx_result


def test_hash_is_stable_but_scoped_to_owner_and_bank() -> None:
    parser = StatementParser(ZoneInfo("Europe/Moscow"))
    content = (FIXTURES / "sber.csv").read_bytes()
    owner = uuid4()
    first = parser.parse(content=content, filename="a.csv", user_id=owner, bank_hint=None)
    again = parser.parse(
        content=content, filename="another.csv", user_id=owner, bank_hint=BankSource.SBER
    )
    other = parser.parse(content=content, filename="a.csv", user_id=uuid4(), bank_hint=None)
    assert first.transactions[0].external_id == again.transactions[0].external_id
    assert first.transactions[0].external_id != other.transactions[0].external_id


def test_comma_csv_with_quoted_decimal() -> None:
    rows = list(csv.reader(io.StringIO((FIXTURES / "sber.csv").read_text()), delimiter=";"))
    stream = io.StringIO()
    csv.writer(stream, delimiter=",").writerows(rows)
    parsed = StatementParser(ZoneInfo("Europe/Moscow")).parse(
        content=stream.getvalue().encode(),
        filename="x.csv",
        user_id=uuid4(),
        bank_hint=None,
    )
    assert parsed.transactions[0].amount == Decimal("1250.50")


@pytest.mark.parametrize(
    ("before", "after"),
    [
        ("-1250.50", "NaN"),
        ("-1250.50", "Infinity"),
        ("-1250.50", "-1.234"),
        ("-1250.50", "-1000000000000"),
        ("5411", "abcd"),
        ("02.10.2026 12:30:00", "31.02.2026 12:30:00"),
        ("RUB", "USD"),
    ],
)
def test_invalid_rows_fail_explicitly(before: str, after: str) -> None:
    content = (FIXTURES / "tbank.csv").read_text().replace(before, after)
    with pytest.raises(StatementError, match="Строка"):
        StatementParser(ZoneInfo("Europe/Moscow")).parse(
            content=content.encode(),
            filename="x.csv",
            user_id=uuid4(),
            bank_hint=None,
        )


def test_ambiguous_sber_direction_is_not_guessed() -> None:
    content = (
        "Дата операции;Время операции;Описание операции;Сумма в валюте счета;Валюта счета;Статус\n"
        "01.10.2026;10:00;Неизвестный мерчант;450,00;RUB;Исполнена\n"
    )
    with pytest.raises(StatementError, match="направление"):
        StatementParser(ZoneInfo("Europe/Moscow")).parse(
            content=content.encode(),
            filename="x.csv",
            user_id=uuid4(),
            bank_hint=None,
        )


@pytest.mark.parametrize(
    ("content", "name"),
    [
        (b"", "a.csv"),
        (b"x", "a.exe"),
        (b"x", "a.xlsx"),
        (b"a;b\n1;2", "a.csv"),
        (b"x" * (MAX_FILE_BYTES + 1), "a.csv"),
    ],
)
def test_bad_files_are_rejected(content: bytes, name: str) -> None:
    with pytest.raises(StatementError):
        StatementParser(ZoneInfo("Europe/Moscow")).parse(
            content=content,
            filename=name,
            user_id=uuid4(),
            bank_hint=None,
        )


def test_more_than_1000_rows_is_not_silently_truncated() -> None:
    lines = (FIXTURES / "tbank.csv").read_text().splitlines()
    content = (lines[0] + "\n" + "\n\n".join([lines[2]] * 1001)).encode()
    with pytest.raises(StatementError, match="1000"):
        StatementParser(ZoneInfo("Europe/Moscow")).parse(
            content=content,
            filename="x.csv",
            user_id=uuid4(),
            bank_hint=None,
        )


def test_bank_hint_must_match_headers() -> None:
    with pytest.raises(StatementError, match="Заголовки"):
        StatementParser(ZoneInfo("Europe/Moscow")).parse(
            content=(FIXTURES / "tbank.csv").read_bytes(),
            filename="x.csv",
            user_id=uuid4(),
            bank_hint=BankSource.SBER,
        )


@pytest.mark.parametrize("problem", ["formula", "too-many-rows", "archive-too-large"])
def test_untrusted_xlsx_is_rejected(problem: str) -> None:
    stream = io.BytesIO()
    if problem == "archive-too-large":
        with ZipFile(stream, "w", compression=ZIP_DEFLATED) as archive:
            archive.writestr("large.xml", b"x" * (31 * 1024 * 1024))
    else:
        workbook = Workbook()
        sheet = workbook.active
        assert sheet is not None
        rows = list(csv.reader(io.StringIO((FIXTURES / "sber.csv").read_text()), delimiter=";"))
        sheet.append(rows[0])
        for _ in range(1001 if problem == "too-many-rows" else 1):
            sheet.append(rows[1])
        if problem == "formula":
            sheet.cell(2, 4).value = "=1+1"
        workbook.save(stream)
    with pytest.raises(StatementError):
        StatementParser(ZoneInfo("Europe/Moscow")).parse(
            content=stream.getvalue(),
            filename="x.xlsx",
            user_id=uuid4(),
            bank_hint=None,
        )


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("RU POS OOO LENTA SPB TERMINAL 1234", "lenta"),
        ("OOO SAMOKAT", "samokat"),
        ("PYATEROCHKA 12344", "pyaterochka"),
        ("Yandex Go", "yandex go"),
        ("Магнит Косметик", "магнит косметик"),
    ],
)
def test_merchant_cleaning(raw: str, expected: str) -> None:
    assert clean_merchant(raw) == expected
