import csv
import hashlib
import io
import json
import re
from datetime import UTC, date, datetime
from decimal import Decimal, InvalidOperation
from itertools import islice
from uuid import UUID
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile
from zoneinfo import ZoneInfo

from defusedxml.common import DefusedXmlException
from pydantic import ValidationError

from src.db.models import BankSource, OperationType
from src.modules.statements.schemas import ParsedStatement, StandardTransactionDTO, StatementError

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 1000
MAX_XLSX_BYTES = 30 * 1024 * 1024

HEADERS = {
    BankSource.TBANK: {"дата операции", "сумма операции", "валюта операции", "описание", "статус"},
    BankSource.SBER: {
        "дата операции",
        "время операции",
        "описание операции",
        "сумма в валюте счета",
        "валюта счета",
        "статус",
    },
    BankSource.VTB: {
        "дата и время операции",
        "сумма операции",
        "валюта операции",
        "тип операции",
        "статус операции",
    },
}


def clean_merchant(description: str) -> str:
    merchant = description.lower().replace("ё", "е")
    merchant = re.sub(r"^(?:(?:ru|rus|retail|card2card|pos|ooo|ооо)\b[\s:*/_-]*)+", "", merchant)
    merchant = re.sub(r"\b(?:терминал|terminal|auth|authorization|rrn|tid)\b.*$", "", merchant)
    merchant = re.sub(r"\b(?:ru|rus|moscow|spb|москва)\b.*$", "", merchant)
    merchant = re.sub(r"[\s#*_]+", " ", merchant)
    merchant = re.sub(r"\s+\d{4,}\s*$", "", merchant)
    return " ".join(merchant.split()).strip(" ,:;-/")


class StatementParser:
    def __init__(self, statement_timezone: ZoneInfo) -> None:
        self._statement_timezone = statement_timezone

    def parse(
        self,
        *,
        content: bytes,
        filename: str,
        user_id: UUID,
        bank_hint: BankSource | None,
    ) -> ParsedStatement:
        rows = self.read_rows(content=content, filename=filename)
        if len(rows) < 2:
            raise StatementError("В выписке отсутствуют операции")
        headers = [value.strip().lower().replace("ё", "е") for value in rows[0]]
        if len(set(headers)) != len(headers) or "" in headers:
            raise StatementError("Заголовки должны быть непустыми и уникальными")
        candidates = [bank for bank, fields in HEADERS.items() if fields <= set(headers)]
        if len(candidates) != 1 or (bank_hint is not None and bank_hint != candidates[0]):
            raise StatementError("Заголовки не соответствуют выбранному или поддерживаемому банку")
        bank = candidates[0]
        if bank is BankSource.VTB and not (
            {"детали операции", "описание", "детали операции / описание"} & set(headers)
        ):
            raise StatementError("Отсутствует колонка описания операции ВТБ")
        transactions = []
        skipped = 0
        for line, values in enumerate(rows[1:], 2):
            if len(values) != len(headers):
                raise StatementError(f"Строка {line}: число ячеек не совпадает с заголовком")
            row = dict(zip(headers, values, strict=True))
            try:
                transaction = self._normalize(row, bank, user_id)
            except (ValueError, InvalidOperation, ValidationError) as err:
                # Do not echo descriptions, card details, or raw workbook values.
                raise StatementError(
                    f"Строка {line}: проверьте дату, сумму, валюту и направление операции"
                ) from err
            if transaction is None:
                skipped += 1
            else:
                transactions.append(transaction)
        return ParsedStatement(
            bank_source=bank,
            total_rows=len(rows) - 1,
            skipped_rows=skipped,
            transactions=transactions,
        )

    def read_rows(self, *, content: bytes, filename: str) -> list[list[str]]:
        if not content or len(content) > MAX_FILE_BYTES:
            raise StatementError("Файл должен быть непустым и не превышать 5 МБ")
        if filename.lower().endswith(".csv"):
            return self._csv_rows(content)
        if filename.lower().endswith(".xlsx"):
            return self._xlsx_rows(content)
        raise StatementError("Поддерживаются только CSV и XLSX")

    def _csv_rows(self, content: bytes) -> list[list[str]]:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = content.decode("cp1251")
            except UnicodeDecodeError as err:
                raise StatementError("Нужна кодировка UTF-8 или Windows-1251") from err
        if "\x00" in text:
            raise StatementError("CSV содержит недопустимые нулевые байты")
        try:
            dialect = csv.Sniffer().sniff(text[:65536], delimiters=";,")
            reader = (
                row
                for row in csv.reader(io.StringIO(text), dialect, strict=True)
                if any(cell.strip() for cell in row)
            )
            rows = list(islice(reader, MAX_ROWS + 2))
        except csv.Error as err:
            raise StatementError("Некорректный CSV или разделитель колонок") from err
        if len(rows) > MAX_ROWS + 1:
            raise StatementError("Допускается не более 1000 строк операций")
        return rows

    def _xlsx_rows(self, content: bytes) -> list[list[str]]:
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException

        try:
            with ZipFile(io.BytesIO(content)) as archive:
                if (
                    len(archive.infolist()) > 200
                    or sum(info.file_size for info in archive.infolist()) > MAX_XLSX_BYTES
                ):
                    raise StatementError("Слишком большой распакованный XLSX")
            workbook = load_workbook(
                io.BytesIO(content), read_only=True, data_only=False, keep_links=False
            )
        except StatementError:
            raise
        except (
            BadZipFile,
            InvalidFileException,
            KeyError,
            OSError,
            ValueError,
            ParseError,
            DefusedXmlException,
        ) as err:
            raise StatementError("Не удалось прочитать XLSX") from err
        try:
            sheet = workbook.active
            if sheet is None:
                raise StatementError("В XLSX нет листа с операциями")
            if sheet.max_row is not None and sheet.max_row > MAX_ROWS + 1:
                raise StatementError("Допускается не более 1000 строк операций")
            if sheet.max_column is not None and sheet.max_column > 32:
                raise StatementError("В XLSX слишком много колонок")
            rows = []
            for cells in islice(sheet.iter_rows(max_col=32), MAX_ROWS + 2):
                if any(cell.data_type == "f" for cell in cells):
                    raise StatementError("Формулы в выписке не поддерживаются")
                values = [
                    cell.value.isoformat()
                    if isinstance(cell.value, (date, datetime))
                    else str(cell.value)
                    if cell.value is not None
                    else ""
                    for cell in cells
                ]
                while values and not values[-1]:
                    values.pop()
                if values:
                    rows.append(values)
            if len(rows) > MAX_ROWS + 1:
                raise StatementError("Допускается не более 1000 строк операций")
            if rows:
                width = len(rows[0])
                for row in rows[1:]:
                    if len(row) < width:
                        row.extend([""] * (width - len(row)))
            return rows
        except StatementError:
            raise
        except (BadZipFile, ParseError, DefusedXmlException, KeyError, ValueError) as err:
            raise StatementError("Некорректная структура XLSX") from err
        finally:
            workbook.close()

    def _normalize(
        self,
        row: dict[str, str],
        bank: BankSource,
        user_id: UUID,
    ) -> StandardTransactionDTO | None:
        status = row.get("статус", row.get("статус операции", "")).strip().lower()
        if status in {
            "failed",
            "в обработке",
            "отклонена",
            "отменена",
            "отклонено",
            "не исполнена",
        }:
            return None
        if status not in {"ok", "исполнена", "исполнено", "успешно"}:
            raise ValueError("Неизвестный статус операции")
        amount_text = row["сумма в валюте счета" if bank is BankSource.SBER else "сумма операции"]
        amount_text = re.sub(r"(?:руб\.?|rub|rur|₽)", "", amount_text, flags=re.IGNORECASE)
        signed_amount = Decimal("".join(amount_text.split()).replace(",", ".").replace("−", "-"))
        if not signed_amount.is_finite():
            raise ValueError("Сумма должна быть конечной")
        if signed_amount == 0:
            return None
        currency = row["валюта счета" if bank is BankSource.SBER else "валюта операции"]
        if currency.strip().lower() not in {"rub", "rur", "руб", "руб.", "₽"}:
            raise ValueError("Поддерживаются только рубли")
        operation_type = self._operation_type(row, bank, signed_amount, amount_text)
        if bank is BankSource.SBER:
            description = row["описание операции"]
            day = row["дата операции"].strip().split("T")[0]
            timestamp = f"{day} {row['время операции'].strip()}"
        elif bank is BankSource.VTB:
            description = row.get(
                "детали операции", row.get("детали операции / описание", row.get("описание", ""))
            )
            timestamp = row["дата и время операции"].strip()
        else:
            description = row["описание"]
            timestamp = row["дата операции"].strip()
        transaction_date = self._timestamp(timestamp)
        merchant = clean_merchant(description)
        amount = abs(signed_amount)
        if amount != amount.quantize(Decimal("0.01")):
            raise ValueError("Допустимо не более двух знаков после запятой")
        mcc_text = row.get("mcc", "").strip()
        if mcc_text and re.fullmatch(r"[0-9]{4}", mcc_text) is None:
            raise ValueError("Некорректный MCC")
        fingerprint = json.dumps(
            [
                str(user_id),
                bank.value,
                transaction_date.isoformat(),
                f"{amount:.2f}",
                operation_type.value,
                merchant,
            ],
            ensure_ascii=False,
            separators=(",", ":"),
        )
        return StandardTransactionDTO(
            external_id=hashlib.sha256(fingerprint.encode()).hexdigest(),
            bank_source=bank,
            transaction_date=transaction_date,
            amount=amount,
            operation_type=operation_type,
            raw_description=description,
            merchant_name_clean=merchant,
            mcc_code=int(mcc_text) if mcc_text else None,
            bank_category_raw=row.get("категория", "").strip() or None,
        )

    def _timestamp(self, value: str) -> datetime:
        try:
            parsed = datetime.fromisoformat(value)
        except ValueError:
            for pattern in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M"):
                try:
                    parsed = datetime.strptime(value, pattern).replace(
                        tzinfo=self._statement_timezone
                    )
                    break
                except ValueError:
                    continue
            else:
                raise ValueError("Некорректная дата") from None
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=self._statement_timezone)
        return parsed.astimezone(UTC)

    def _operation_type(
        self,
        row: dict[str, str],
        bank: BankSource,
        signed_amount: Decimal,
        amount_text: str,
    ) -> OperationType:
        if bank is BankSource.TBANK:
            return OperationType.DEBIT if signed_amount < 0 else OperationType.CREDIT
        marker = row.get("тип операции", "").strip().lower()
        if marker in {"списание", "расход", "оплата", "debit"}:
            return OperationType.DEBIT
        if marker in {"пополнение", "зачисление", "доход", "credit"}:
            return OperationType.CREDIT
        if bank is BankSource.SBER:
            if signed_amount < 0:
                return OperationType.DEBIT
            if amount_text.strip().startswith("+"):
                return OperationType.CREDIT
            description = row["описание операции"].lower()
            if re.search(r"\b(списание|оплата|покупка)\b", description):
                return OperationType.DEBIT
            if re.search(r"\b(зачисление|пополнение|зарплата)\b", description):
                return OperationType.CREDIT
        raise ValueError("Невозможно определить направление операции")
